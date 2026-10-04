import docx
import sys
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('NEXUS_Complete_Technical_Documentation.docx')

print("=" * 80)
print("NEXUS TECHNICAL DOCUMENTATION VERIFICATION AUDIT")
print("=" * 80)

# Section dimensions
sec = doc.sections[0]
print(f"Page Dimensions: {sec.page_width.inches:.2f} in x {sec.page_height.inches:.2f} in (Target A4: 8.27 in x 11.69 in)")
print(f"Margins: Left={sec.left_margin.inches:.2f} in, Right={sec.right_margin.inches:.2f} in, Top={sec.top_margin.inches:.2f} in, Bottom={sec.bottom_margin.inches:.2f} in")

# Body paragraphs alignment
body_ps = [
    p for p in doc.paragraphs
    if not p.style.name.startswith('Heading')
    and not p.style.name.startswith('List')
    and len(p.text) > 40
    and not p.text.startswith('Figure ')
    and not p.text.startswith('Listing:')
    and not p.text.startswith('FINAL YEAR')
    and not p.text.startswith('AI-Powered')
    and not p.text.startswith('A Closed-Loop')
    and not p.text.startswith('Table ')
]

print(f"\nSubstantive Body Paragraphs (>40 chars): {len(body_ps)}")
justified_count = sum(1 for p in body_ps if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY)
print(f"Justified Alignment Count: {justified_count} / {len(body_ps)} ({justified_count/len(body_ps)*100:.1f}%)")

# Headings
h1s = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
h2s = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
h3s = [p.text for p in doc.paragraphs if p.style.name == 'Heading 3']
print(f"\nHeadings Hierarchy:")
print(f"  Heading 1 Sections: {len(h1s)}")
print(f"  Heading 2 Subsections: {len(h2s)}")
print(f"  Heading 3 Subsections: {len(h3s)}")

# Viva questions
viva_qs = [p.text for p in doc.paragraphs if p.text.startswith('Q') and ':' in p.text[:5]]
print(f"\nTotal Viva Questions: {len(viva_qs)} across 8 categories")

# Figure Captions
fig_caps = [p.text for p in doc.paragraphs if p.text.startswith('Figure ')]
print(f"\nTotal Figure Captions ({len(fig_caps)}):")
for fc in fig_caps[:6]:
    print('  ', fc)
print(f'   ... plus {len(fig_caps) - 6} more figures')

# Table Titles
tbl_titles = [p.text for p in doc.paragraphs if p.text.startswith('Table ')]
print(f"\nTotal Formally Numbered Table Captions ({len(tbl_titles)}):")
for tt in tbl_titles:
    print('  ', tt)

# Total Words
words = sum(len(p.text.split()) for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            words += sum(len(p.text.split()) for p in cell.paragraphs)
print(f"\nTotal Word Count: {words:,} words")
print(f"Estimated Standard A4 Academic Pages: ~{int(words / 320)} - {int(words / 270)} pages")
print("=" * 80)
