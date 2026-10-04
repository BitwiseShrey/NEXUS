"""
Audit script for NEXUS_Project_Review_Pastel_13_Slides.pptx and .pdf
"""

import os
from pathlib import Path
from pptx import Presentation
import fitz

ROOT_DIR = Path(__file__).resolve().parent.parent
PPTX_PATH = ROOT_DIR / "NEXUS_Project_Review_Pastel_13_Slides.pptx"
PDF_PATH = ROOT_DIR / "NEXUS_Project_Review_Pastel_13_Slides.pdf"
AUDIT_IMG_DIR = ROOT_DIR / "NEXUS_Pastel_Presentation_Assets" / "audit_slides"
AUDIT_IMG_DIR.mkdir(parents=True, exist_ok=True)

print("=" * 70)
print(" AUDITING 13-SLIDE PASTEL PPTX PRESENTATION ")
print("=" * 70)

prs = Presentation(str(PPTX_PATH))
print(f"Total Slides in PPTX: {len(prs.slides)}")
assert len(prs.slides) == 13, f"Expected 13 slides, got {len(prs.slides)}"

for idx, slide in enumerate(prs.slides, 1):
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    shapes_count = len(slide.shapes)
    words = len(notes.split())
    print(f"Slide {idx:02d}: Shapes={shapes_count:02d} | Notes Words={words:03d} | Notes Preview: {notes[:65].strip()}...")
    assert len(notes.strip()) > 40, f"Slide {idx} missing speaker notes!"

print("\nAll 13 PPTX slides verified with complete speaker notes!")

print("\n" + "=" * 70)
print(" AUDITING 13-PAGE PASTEL PDF PRESENTATION ")
print("=" * 70)

doc = fitz.open(str(PDF_PATH))
print(f"Total Pages in PDF: {len(doc)}")
assert len(doc) == 13, f"Expected 13 PDF pages, got {len(doc)}"

for p_num in range(len(doc)):
    page = doc[p_num]
    pix = page.get_pixmap(dpi=150)
    out_path = AUDIT_IMG_DIR / f"pastel_slide_{p_num+1:02d}.png"
    pix.save(str(out_path))

print(f"Rendered {len(doc)} slide images to: {AUDIT_IMG_DIR}")
print("Pastel Presentation Audit Passed with 100% Verification!")
