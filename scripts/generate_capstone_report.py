"""
NEXUS - DSN4091 Capstone Project Phase-I Report Generator
Generates an official Capstone Report following the exact structure and formatting
of the VIT Bhopal University Capstone Phase-I template.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS_DIR = os.path.join(ROOT_DIR, "NEXUS_Documentation_Assets")
OUTPUT_DOCX = os.path.join(ROOT_DIR, "NEXUS_Capstone_Phase_I_Report.docx")

# Styling Constants
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_DARK = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)
HEX_LIGHT_GRAY = "F8FAFC"
HEX_BORDER = "CCCCCC"


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, border_hex="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def init_report_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Footer: Page | X
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.space_before = Pt(8)
        frun = fp.add_run("Page | ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(10)
        frun.font.color.rgb = COLOR_BLACK

        fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        fp._p.append(fldSimple)

    return doc


def add_para(doc, text="", bold=False, italic=False, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.font.color.rgb = COLOR_BLACK
    return p


def add_bullet(doc, bold_prefix, text, space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    r_pre = p.add_run(bold_prefix)
    r_pre.bold = True
    r_pre.font.name = "Times New Roman"
    r_pre.font.size = Pt(11.5)
    r_pre.font.color.rgb = COLOR_BLACK

    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(11.5)
    r_txt.font.color.rgb = COLOR_BLACK
    return p


def add_numbered_item(doc, number_str, text, bold=False, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    r_num = p.add_run(number_str)
    r_num.bold = True
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11.5)

    r_txt = p.add_run(text)
    r_txt.bold = bold
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(11.5)
    return p


def add_subitem(doc, label, text, space_after=3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run(label + " ")
    r_lbl.bold = True
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)

    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(11)
    return p


def add_chapter_heading(doc, chapter_num, title):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(18)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(f"CHAPTER {chapter_num}")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(16)
    r1.font.color.rgb = COLOR_BLACK

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run(title)
    r2.bold = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.color.rgb = COLOR_BLACK


def add_section_heading(doc, sec_num, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"{sec_num}  {title}")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_BLACK


def add_subsection_heading(doc, subsec_num, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{subsec_num} {title}")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_BLACK


def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, "1E293B")
    set_cell_margins(cell, 80, 80, 100, 100)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05

    for line in code_str.strip().split("\n"):
        r = p.add_run(line + "\n")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(226, 232, 240)


def add_figure_image(doc, img_name, fig_label, caption, width_in=5.8):
    img_path = os.path.join(ASSETS_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Fig {fig_label}: {caption}")
        r_cap.bold = True
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.color.rgb = RGBColor(180, 40, 40)
    else:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run(f"[Fig {fig_label}: {caption}]")
        r_cap.italic = True
        r_cap.font.size = Pt(10)


print("Loaded report builder module.")
