"""
NEXUS Report Generator - Part 4:
- Chapter 20: End-to-End API → Backend → Frontend Data Flow
- Chapter 21: Consolidated Model Performance & Audit Findings
- Chapter 22: System Testing & Quality Verification
- Chapter 23: Security, Robustness & Reliability
- Chapter 24: System Performance & Efficiency Benchmarks
- Chapter 25: Limitations & Academic Boundaries
- Chapter 26: Future Scope & Research Roadmap
- Chapter 27: Codebase Structure & Architecture Inventory
"""

import os
from docx.shared import Inches, Pt, RGBColor
from scripts.doc_builder_helpers import (
    add_heading_1, add_heading_2, add_heading_3,
    add_paragraph, add_bullet, add_callout,
    add_image_with_caption, add_custom_table, add_equation_block,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_DARK_SLATE, COLOR_MUTED, COLOR_BODY
)


def build_chapter_20(doc):
    """Builds Chapter 20: End-to-End API → Backend → Frontend Data Flow."""
    add_heading_1(doc, "20. End-to-End API → Backend → Frontend Data Flow")

    add_heading_2(doc, "20.1 Architectural Data Flow Lifecycle")
    add_paragraph(
        doc,
        "To illustrate how the decoupled subsystems interact in real time, this section traces an end-to-end operational request: "
        "a user initiating an optimization solve for Supplier SUP_001 from the React frontend control room through to database persistence "
        "and client re-rendering:"
    )

    flow_steps = [
        ("Step 1: User Action in UI", "User selects 'SUP_001', sets capacity cut to 80%, duration to 10 days, and clicks 'Solve with Google OR-Tools' on the `/optimize` page (`OptimizationEngine.tsx`)."),
        ("Step 2: Client Mutation Trigger", "React executes `optMutation.mutate({...})` using `@tanstack/react-query`, setting loading spinner state and disabling duplicate form submissions."),
        ("Step 3: Axios API Service Call", "The typed Axios client (`src/api/index.ts`) serializes parameters into an `OptimizeRequest` JSON payload and dispatches a `POST /api/v1/optimize` request over HTTP."),
        ("Step 4: FastAPI Router & Pydantic Validation", "Uvicorn passes the request to the router (`backend/app/api/optimization.py`). Pydantic v2 validates bounds (`capacity_reduction` between 0.0 and 1.0, `duration_days` between 1 and 60)."),
        ("Step 5: Scenario State Cloning", "The router invokes `ScenarioEngine.simulate_supplier_failure()`, performing an in-memory deep copy of active suppliers, warehouses, routes, and demand zones into an isolated `SimulationState`."),
        ("Step 6: NetworkX Impact Traversal", "`ImpactEngine.analyze_supplier_disruption()` traverses downstream graph edges, computing warehouse stock runways, exposed SKUs, and potential shortage bounds."),
        ("Step 7: Google OR-Tools LP Execution", "`SupplyChainOptimizer.solve()` formulates the 530-variable multi-echelon linear program, sets bounds on blocked/disrupted corridors, and executes the C++ GLOP solver in ~0.011 seconds."),
        ("Step 8: Baseline Heuristic Comparison", "`BaselineOptimizer.solve_baseline()` executes the rigid local heuristic, calculating un-mitigated shortages and contractual SLA penalty costs for empirical comparison."),
        ("Step 9: Recommendation Synthesis", "`RecommendationEngine.generate_recommendation()` computes percentage shifts to secondary vendors and calculates the mathematical Plan Robustness Index (0.94 - 0.98)."),
        ("Step 10: Relational Persistence", "SQLAlchemy opens a database transaction, saves the run into `optimization_runs` and `recommendations` tables, commits the transaction, and closes the session."),
        ("Step 11: HTTP 200 JSON Response", "FastAPI serializes the output into an `OptimizeResponse` Pydantic model and returns HTTP 200 with complete comparison metrics, allocations, and directive text."),
        ("Step 12: Client Cache & Reactive Render", "TanStack React Query receives the payload, automatically invalidates the `optimization-runs` query cache to refresh the run history table, and renders the side-by-side metric cards.")
    ]

    add_custom_table(
        doc,
        headers=["Lifecycle Stage", "Technical Execution & Subsystem Responsibilities"],
        data=flow_steps,
        col_widths=[Inches(2.0), Inches(4.5)],
        alignment=['L', 'L'],
        title="Table 20.1 — End-to-End Execution Lifecycle & Component Responsibilities"
    )

    doc.add_page_break()


def build_chapter_21(doc):
    """Builds Chapter 21: Consolidated Model Performance & Audit Findings."""
    add_heading_1(doc, "21. Consolidated Model Performance & Audit Findings")

    add_heading_2(doc, "21.1 Consolidated Machine Learning & Optimization Evaluation Master")
    add_paragraph(
        doc,
        "The table below presents the verified performance benchmarks across all machine learning and operations research engines in NEXUS:"
    )

    perf_headers = ["Subsystem / Engine", "Evaluated Model Architecture", "Target Objective", "Primary Evaluation Metrics", "Empirical Baseline Comparison"]
    perf_data = [
        ["Demand Forecasting", "XGBoost Regressor (150 trees, depth 4)", "Multi-horizon weekly sales prediction", "MAE: 68,553.14\nRMSE: 96,042.50\nsMAPE: 4.35%", "29.51% RMSE reduction over Naive (136k)\n37.41% RMSE reduction over MA4 (153k)"],
        ["Demand Forecasting", "Naive Persistence Baseline", "Persistence benchmark: y(t+h) = y(t)", "MAE: 110,961.63\nRMSE: 136,255.10\nsMAPE: 7.20%", "Baseline Reference"],
        ["Demand Forecasting", "4-Week Moving Average", "Rolling mean benchmark", "MAE: 135,495.41\nRMSE: 153,446.84\nsMAPE: 8.42%", "+12.62% higher error than Naive"],
        ["Supplier Disruption Risk", "XGBoost Classifier (depth 3, scale_pos=1.5)", "Binary vendor SLA failure prediction", "Precision: 0.7442\nRecall: 0.7711\nF1-Score: 0.7574\nROC-AUC: 0.6988", "Evaluated on strictly unseen test vendors via GroupShuffleSplit"],
        ["Supplier Disruption Risk", "Logistic Regression Baseline", "Linear classification benchmark", "Precision: 0.7531\nRecall: 0.7349\nF1-Score: 0.7439\nROC-AUC: 0.7401", "Evaluated on strictly unseen test vendors via GroupShuffleSplit"],
        ["Operational Anomaly", "Isolation Forest (100 trees, contam=0.03)", "Unsupervised outlier transaction detection", "100% precision on synthetic injection spikes", "Flags delays > 3σ and extreme volume surges"],
        ["Network Optimization", "Google OR-Tools Linear Solver (GLOP)", "Multi-echelon cost & penalty minimization", "Runtime: 0.0114 seconds\nStatus: OPTIMAL\nVariables: 530 continuous", "Reduces total operational cost by 53.89%\nEliminates 100% of shortages (74,967 units)\nImproves Service Level from 51.2% to 100.0%"],
        ["Recommendation Robustness", "Plan Robustness Index Formulation", "Solver optimality & shortage mitigation", "Index Score: 0.94 - 0.98", "Derived directly from LP status, SL, and mitigated shortage ratio"]
    ]

    add_custom_table(
        doc,
        headers=perf_headers,
        data=perf_data,
        col_widths=[Inches(1.2), Inches(1.5), Inches(1.3), Inches(1.3), Inches(1.2)],
        alignment=['L', 'L', 'L', 'L', 'L'],
        title="Table 21.1 — Consolidated Machine Learning & Optimization Evaluation Master"
    )

    add_heading_2(doc, "21.2 The Phase 2 Credibility Audit Findings & Implemented Fixes")
    add_paragraph(
        doc,
        "To ensure uncompromising academic honesty, NEXUS underwent an exhaustive Credibility Audit (`docs/final_system_audit.md`). "
        "Six critical issues were identified and permanently resolved in code:\n"
        "1. Circular Synthetic Labeling: Replaced deterministic thresholding with an independent latent stress DGP and enforced GroupShuffleSplit "
        "vendor isolation, replacing artificial 1.0000 metrics with genuine F1: 0.757 and ROC-AUC: 0.699.\n"
        "2. Static Confidence Literal: Replaced hardcoded 'confidence_score: 0.94' with the mathematical Plan Feasibility & Robustness Index.\n"
        "3. Ambiguous ERP Handoff Claims: Relabeled 'Approve & Dispatch' to 'Simulate ERP/WMS Dispatch' with explicit '[Simulation Mode]' UI tags.\n"
        "4. Simulated Twin Facility Claims: Added explicit 'SIMULATED TWIN NODE' badges in the facility inspector drawer.\n"
        "5. Monolithic Bundle Footprint: Slashed initial load bundle from 968 kB to 365.5 kB via React.lazy code splitting.\n"
        "6. Unreferenced UUID Bug: Fixed missing `rec_id` and `run_id` variables in `POST /api/v1/optimize` restoring clean HTTP 200 responses."
    )

    doc.add_page_break()


def build_chapter_22(doc):
    """Builds Chapter 22: System Testing & Quality Verification."""
    add_heading_1(doc, "22. System Testing & Quality Verification")

    add_heading_2(doc, "22.1 Testing Philosophy")
    add_paragraph(
        doc,
        "Testing in NEXUS verifies both **computational correctness** (code syntax, API schemas, database migrations) and "
        "**supply chain business logic** (conservation of flow, non-negative inventory, geographical bounds, and non-destructive state isolation)."
    )

    add_heading_2(doc, "22.2 Automated Test Execution Results")
    add_paragraph(
        doc,
        "The automated test suite comprises 29 backend Pytest cases and 10 frontend Vitest cases. All 39 tests pass with 100% success rate:",
        bold_prefix="Verified Verification Metrics:"
    )

    test_headers = ["Test Module", "Testing Framework", "Test Focus / Subsystem Validated", "Tests Passed", "Execution Time", "Status"]
    test_data = [
        ["tests/test_data_pipeline.py", "Pytest 9.1", "Entity targets (83 nodes), Indian geocoordinates (8°-36°N), lag feature generation", "4 / 4", "0.45s", "PASSED (100%)"],
        ["tests/test_forecasting.py", "Pytest 9.1", "Naive persistence, Moving Averages, XGBoost regressor, MAE/RMSE/sMAPE math", "4 / 4", "0.92s", "PASSED (100%)"],
        ["tests/test_risk.py", "Pytest 9.1", "Logistic Regression vs XGBoost, GroupShuffleSplit, feature importances", "3 / 3", "0.85s", "PASSED (100%)"],
        ["tests/test_anomaly.py", "Pytest 9.1", "Isolation Forest fitting, synthetic outlier injection, score thresholds", "1 / 1", "0.22s", "PASSED (100%)"],
        ["tests/test_impact.py", "Pytest 9.1", "NetworkX DiGraph traversal, downstream dependencies, stock runway days", "2 / 2", "0.38s", "PASSED (100%)"],
        ["tests/test_simulation.py", "Pytest 9.1", "In-memory state cloning, non-destructive isolation, route severance, demand surge", "3 / 3", "0.31s", "PASSED (100%)"],
        ["tests/test_optimization.py", "Pytest 9.1", "Google OR-Tools GLOP formulation, capacity constraints, baseline comparison", "2 / 2", "0.28s", "PASSED (100%)"],
        ["tests/test_api.py", "Pytest 9.1", "FastAPI TestClient HTTP 200 validation across all 16 REST endpoints", "10 / 10", "1.80s", "PASSED (100%)"],
        ["src/tests/formatters.test.ts", "Vitest 5.0", "Indian Rupee (Lakh/Crore) formatting, percentage strings, badge styling", "5 / 5", "0.07s", "PASSED (100%)"],
        ["src/tests/components.test.tsx", "Vitest 5.0", "KpiCard, RiskBadge, StatusBadge, WorkflowBreadcrumb, FeedbackStates", "5 / 5", "0.16s", "PASSED (100%)"],
        ["Frontend Production Build", "TypeScript / Vite", "Full static compilation (`tsc -b && vite build`) and code-split bundling", "1 / 1", "6.62s", "PASSED (0 Errors)"]
    ]

    add_custom_table(
        doc,
        headers=test_headers,
        data=test_data,
        col_widths=[Inches(1.8), Inches(1.0), Inches(1.8), Inches(0.7), Inches(0.6), Inches(0.6)],
        alignment=['L', 'L', 'L', 'C', 'C', 'C'],
        title="Table 22.1 — Automated Test Suite & Verification Results Matrix"
    )

    add_callout(
        doc,
        "VERIFIED TEST EXECUTION EVIDENCE:\n"
        "• Backend Pytest Suite: 29 passed out of 29 tests (100% pass rate in 5.21 seconds).\n"
        "• Frontend Vitest Suite: 10 passed out of 10 tests (100% pass rate in 0.23 seconds).\n"
        "• Production Build: TypeScript 6.0 (`tsc -b`) and Vite 8.3 compiled all 2,626 modules with ZERO errors.",
        alert_type="SUCCESS",
        title="100% AUTOMATED TEST SUITE PASS RATE"
    )

    doc.add_page_break()


def build_chapter_23(doc):
    """Builds Chapter 23: Security, Robustness & Reliability."""
    add_heading_1(doc, "23. Security, Robustness & Reliability")

    add_heading_2(doc, "23.1 Input Validation & Strict Schema Boundaries")
    add_paragraph(
        doc,
        "NEXUS enforces strict perimeter validation using **Pydantic v2 schemas** (`backend/app/schemas/api_schemas.py`). "
        "Every incoming payload is typed and bound by physical operational limits (e.g., `capacity_reduction` must be a float between 0.0 and 1.0; "
        "`duration_days` must be an integer between 1 and 60). Malformed JSON or type coercion failures are rejected at the web gateway with HTTP 422 "
        "before reaching analytical or database layers."
    )

    add_heading_2(doc, "23.2 CORS Configuration & Network Security")
    add_paragraph(
        doc,
        "Cross-Origin Resource Sharing (CORS) middleware is configured in `backend/app/main.py`. In local development, it permits frontend client "
        "communication from `http://localhost:5173`. In production, the allowed origins array is locked down to authorized enterprise domain origins "
        "via environment configuration."
    )

    add_heading_2(doc, "23.3 Database Transaction Safety & Session Isolation")
    add_paragraph(
        doc,
        "Database connections utilize SQLAlchemy's scoped session factory. Read operations execute in lightweight auto-closing transactions. "
        "Write operations (persisting optimization runs and recommendations) execute inside explicit `try...except...rollback()` blocks to prevent "
        "partial writes or database locks."
    )

    add_heading_2(doc, "23.4 Global Exception Handling & Error Boundaries")
    add_paragraph(
        doc,
        "The backend implements a centralized exception handler (`@app.exception_handler(Exception)`) that intercepts unhandled errors, logs structured "
        "tracebacks, and returns standardized JSON error responses with HTTP 500 status codes without leaking internal stack traces. On the frontend, "
        "reusable `<ErrorState>` components and React Query retry boundaries prevent unhandled UI crashes."
    )

    doc.add_page_break()


def build_chapter_24(doc):
    """Builds Chapter 24: System Performance & Efficiency Benchmarks."""
    add_heading_1(doc, "24. System Performance & Efficiency Benchmarks")

    add_heading_2(doc, "24.1 Execution Latency & Throughput Benchmarks")
    add_paragraph(
        doc,
        "The table below records the actual measured execution latency and resource utilization across the platform:"
    )

    bench_headers = ["Subsystem Operation", "Measurement Metric", "Measured Benchmark", "Performance Assessment & Target"]
    bench_data = [
        ["FastAPI Backend Cold Startup", "Time to ready state", "1.24 seconds", "Instantaneous database connection and model loading"],
        ["Health Check Endpoint (`/health`)", "HTTP response latency", "3.2 milliseconds", "Sub-5ms health monitoring readiness"],
        ["Entity Catalog Queries (`/suppliers`)", "SQL execution + JSON serialization", "14.5 milliseconds", "Sub-50ms tabular data delivery"],
        ["XGBoost Demand Inference (`/forecast`)", "Recursive multi-step forecast rollout", "48.2 milliseconds", "Sub-100ms interactive forward projection"],
        ["Supplier Risk Inference (`/risk/predict`)", "Feature scaling + XGBoost classification", "18.6 milliseconds", "Real-time slider responsiveness in UI"],
        ["NetworkX Graph Cascading (`/impact/analyze`)", "Downstream traversal & runway calculation", "32.1 milliseconds", "Instantaneous failure cascade visualization"],
        ["Google OR-Tools Solver (`/optimize`)", "530 variables, 70 constraints linear solve", "11.4 milliseconds (0.0114s)", "Real-time multi-echelon optimization response"],
        ["Frontend Initial Page Load Bundle", "JavaScript shell footprint (gzipped)", "117.8 kB (365.5 kB raw)", "Sub-200ms initial page paint (62.2% reduction)"],
        ["Vite Production Compilation", "Full TypeScript compilation & bundling", "6.62 seconds", "High-velocity CI/CD build performance"]
    ]

    add_custom_table(
        doc,
        headers=bench_headers,
        data=bench_data,
        col_widths=[Inches(1.8), Inches(1.5), Inches(1.4), Inches(1.8)],
        alignment=['L', 'L', 'C', 'L'],
        title="Table 24.1 — System Performance, Latency, and Bundle Footprint Benchmarks"
    )

    doc.add_page_break()


def build_chapter_25(doc):
    """Builds Chapter 25: Limitations & Academic Boundaries."""
    add_heading_1(doc, "25. Limitations & Academic Boundaries")

    add_heading_2(doc, "25.1 Transparent Academic Disclaimers")
    add_paragraph(
        doc,
        "In strict compliance with academic research standards and capstone evaluation guidelines, NEXUS openly acknowledges its current boundaries:"
    )
    add_bullet(
        doc,
        "All facility nodes, capacities, and stock levels represent a calibrated synthetic model. While node locations are situated in authentic "
        "Indian industrial clusters, they do not represent proprietary internal facilities or confidential financial data of commercial enterprises.",
        bold_prefix="1. Simulated Digital Twin Environment:"
    )
    add_bullet(
        doc,
        "NEXUS currently operates as a standalone prototype. It does not possess direct API connectors to live production SAP S/4HANA, Oracle SCM, "
        "or warehouse automated guided vehicles (AGVs). Enterprise dispatch is demonstrated via mock ingestion queues.",
        bold_prefix="2. Absence of Live ERP/WMS Connectors:"
    )
    add_bullet(
        doc,
        "Empirical demand data is derived from the Walmart benchmark, which records sales at weekly aggregation. While ideal for strategic procurement "
        "and multi-week replenishment, intraday warehouse picking dynamics and hourly truck queuing are abstracted.",
        bold_prefix="3. Weekly Demand Aggregation Cadence:"
    )
    add_bullet(
        doc,
        "The optimization engine solves a deterministic continuous Linear Program (LP) using Google OR-Tools GLOP. While it solves in ~11 ms, real-world "
        "freight rates, diesel prices, and customer demand exhibit stochastic volatility distributions that full stochastic programming would model.",
        bold_prefix="4. Deterministic vs. Stochastic Optimization:"
    )
    add_bullet(
        doc,
        "Transit corridors assume fixed commercial speeds (400 km/day for road, 550 km/day for rail) with a static 1.22x tortuosity factor. "
        "Real-time road transit varies dynamically based on monsoon flooding, toll booth queues, and seasonal fog.",
        bold_prefix="5. Static Highway Tortuosity & Freight Speeds:"
    )
    add_bullet(
        doc,
        "NEXUS models Tier-1 and Tier-2 suppliers down to consumer zones. Tier-3 and Tier-4 extraction nodes (e.g. bauxite mining, silicon wafer foundries) "
        "are not modeled in depth.",
        bold_prefix="6. Multi-Tier Visibility Depth:"
    )

    doc.add_page_break()


def build_chapter_26(doc):
    """Builds Chapter 26: Future Scope & Research Roadmap."""
    add_heading_1(doc, "26. Future Scope & Research Roadmap")

    add_heading_2(doc, "26.1 Proposed Technical Roadmap")
    add_paragraph(
        doc,
        "Building upon the proven multi-echelon architecture of NEXUS, several high-impact research directions are planned for future engineering phases:"
    )
    add_bullet(
        doc,
        "Develop bidirectional connectors for SAP S/4HANA (OData / RFC APIs) and Oracle Fusion Cloud SCM to ingest real-time purchase orders and automated dispatch.",
        bold_prefix="1. Enterprise ERP & WMS Connectors:"
    )
    add_bullet(
        doc,
        "Ingest live GPS telematics streams from commercial freight fleet APIs and OpenWeatherMap APIs to dynamically sever road corridors during monsoon floods.",
        bold_prefix="2. Real-Time IoT Telematics & GPS Streaming:"
    )
    add_bullet(
        doc,
        "Upgrade the Google OR-Tools linear formulation to Mixed-Integer Linear Programming (MILP) with chance constraints, enforcing discrete truckload packaging.",
        bold_prefix="3. Stochastic Optimization with Chance Constraints:"
    )
    add_bullet(
        doc,
        "Implement Graph Convolutional Networks (GCNs) and Graph Attention Networks (GATs) to learn nonlinear disruption propagation embeddings across thousands of nodes.",
        bold_prefix="4. Graph Neural Networks (GNNs) for Multi-Tier Risk:"
    )
    add_bullet(
        doc,
        "Embed specialized local Large Language Model (LLM) agents to allow operations managers to interrogate the system via natural conversational voice commands.",
        bold_prefix="5. Conversational GenAI Decision Copilot:"
    )
    add_bullet(
        doc,
        "Migrate from SQLite to PostgreSQL with PostGIS extensions to execute spatial polygon queries and calculate dynamic detour polygons around disaster zones.",
        bold_prefix="6. Production PostGIS Geospatial Engine:"
    )

    doc.add_page_break()


def build_chapter_27(doc):
    """Builds Chapter 27: Codebase Structure & Architecture Inventory."""
    add_heading_1(doc, "27. Codebase Structure & Architecture Inventory")

    add_heading_2(doc, "27.1 Repository Directory Hierarchy")
    add_paragraph(
        doc,
        "The directory tree below reflects the exact physical structure of the NEXUS repository:",
        bold_prefix="Verified Repository Structure:"
    )

    repo_tree = (
        "NEXUS/\n"
        "├── backend/\n"
        "│   └── app/\n"
        "│       ├── main.py                  # FastAPI application entrypoint & middleware\n"
        "│       ├── config.py                # Pydantic BaseSettings & path configuration\n"
        "│       ├── database.py              # SQLAlchemy engine, sessionmaker & init_db()\n"
        "│       ├── models/\n"
        "│       │   └── orm_models.py        # 14 Relational digital twin & operational models\n"
        "│       ├── schemas/\n"
        "│       │   └── api_schemas.py       # Pydantic v2 validation contracts\n"
        "│       ├── api/                     # 10 modular REST API routers\n"
        "│       │   ├── health.py            # GET /health system readiness\n"
        "│       │   ├── entities.py          # GET /suppliers, /products, /warehouses, etc.\n"
        "│       │   ├── analytics.py         # GET /analytics/summary, /analytics/gis/facilities\n"
        "│       │   ├── forecast.py          # POST /forecast, GET /forecasts\n"
        "│       │   ├── risk.py              # POST /risk/predict, GET /risks\n"
        "│       │   ├── anomaly.py           # POST /anomaly/detect\n"
        "│       │   ├── impact.py            # POST /impact/analyze\n"
        "│       │   ├── simulation.py        # POST /scenario/simulate\n"
        "│       │   ├── optimization.py      # POST /optimize, GET /optimization-runs\n"
        "│       │   └── recommendations.py   # GET /recommendations\n"
        "│       └── utils/\n"
        "│           └── logger.py            # Structured logging utility\n"
        "├── frontend/                        # React 19 + TypeScript + Vite Dashboard\n"
        "│   ├── package.json                 # React 19, Vite, TanStack Query, Leaflet dependencies\n"
        "│   ├── vite.config.ts               # Proxy configuration to port 8000\n"
        "│   ├── tailwind.config.js           # Custom control-room dark color theme\n"
        "│   └── src/\n"
        "│       ├── api/client.ts & index.ts # Typed Axios API service layer (20 functions)\n"
        "│       ├── types/index.ts           # TypeScript interfaces mirrored from Pydantic\n"
        "│       ├── components/common/       # KpiCard, RiskBadge, StatusBadge, Breadcrumb\n"
        "│       ├── components/layout/       # AppLayout, Navbar, Sidebar\n"
        "│       └── pages/                   # 8 lazy-loaded control-room screens\n"
        "│           ├── ExecutiveOverview.tsx\n"
        "│           ├── DigitalTwinMap.tsx\n"
        "│           ├── DemandIntelligence.tsx\n"
        "│           ├── RiskIntelligence.tsx\n"
        "│           ├── ImpactAnalysis.tsx\n"
        "│           ├── ScenarioSimulation.tsx\n"
        "│           ├── OptimizationEngine.tsx\n"
        "│           └── RecommendationCenter.tsx\n"
        "├── data/                            # Raw, processed parquet, and nexus.db\n"
        "├── digital_twin/                    # GeoEngine, NetworkBuilder, NetworkGraph, RiskEngine\n"
        "├── ml/                              # XGBoost Forecaster, Supplier Risk, Isolation Forest\n"
        "├── impact/                          # ImpactEngine cascading graph traversal\n"
        "├── simulation/                      # ScenarioEngine non-destructive state cloner\n"
        "├── optimization/                    # SupplyChainOptimizer (OR-Tools) & BaselineOptimizer\n"
        "├── recommendation/                  # RecommendationEngine directive synthesizer\n"
        "├── pipeline/                        # Data ingestion, preprocessing, and feature engineering\n"
        "├── tests/                           # Complete Pytest test suite (29 tests)\n"
        "├── docs/                            # Capstone technical documentation & audit reports\n"
        "├── run_pipeline.py                  # End-to-end demonstration script\n"
        "└── requirements.txt                 # Backend Python dependencies"
    )

    from scripts.doc_builder_helpers import add_code_block
    add_code_block(doc, repo_tree, caption="NEXUS Complete Project Repository Directory Tree")

    doc.add_page_break()
