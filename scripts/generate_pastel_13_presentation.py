"""
NEXUS — 13-Slide Pastel Presentation Generator
Theme: Pastel Supply Chain Control Tower
Modern SaaS product presentation x Supply-chain network visualization x Clean academic engineering review
Palette: Warm Ivory (#FAF8F3), Soft Cream (#FFFDF9), Mint (#D9F3EE), Blue (#DCEBFA),
         Lavender (#E9E1F7), Peach (#FCE3D6), Yellow (#FFF0C7)
Text: Charcoal (#20252B), Slate (#626A73), Muted (#9299A1)
Presenter: Shreyansh Uttam | B.Tech CSE (AI & ML) | VIT Bhopal University
"""

import os
import sys
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT_DIR = Path(__file__).resolve().parent.parent
ASSETS_DIR = ROOT_DIR / "NEXUS_Pastel_Presentation_Assets"
DOC_ASSETS_DIR = ROOT_DIR / "NEXUS_Documentation_Assets"
PPTX_OUTPUT = ROOT_DIR / "NEXUS_Project_Review_Pastel_13_Slides.pptx"
PDF_OUTPUT = ROOT_DIR / "NEXUS_Project_Review_Pastel_13_Slides.pdf"
SOURCE_COPY = ROOT_DIR / "NEXUS_Presentation_Source" / "generate_pastel_13_presentation.py"

# --- PASTEL COLOR PALETTE DEFINITIONS ---
IVORY = RGBColor(250, 248, 243)       # #FAF8F3 Warm Ivory Canvas
CREAM = RGBColor(255, 253, 249)       # #FFFDF9 Soft Cream Card Fill
BORDER_LIGHT = RGBColor(226, 221, 210) # #E2DDD2 Subtle Card Border

# Charcoal text hierarchy
CHARCOAL = RGBColor(32, 37, 43)       # #20252B Primary Text
SLATE = RGBColor(98, 106, 115)        # #626A73 Secondary Text
MUTED = RGBColor(146, 153, 161)       # #9299A1 Muted / Microcopy

# Pastel Accents & Borders
MINT_BG = RGBColor(217, 243, 238)     # #D9F3EE Mint Fill
MINT_LINE = RGBColor(16, 185, 129)    # #10B981 Mint Border / Text
BLUE_BG = RGBColor(220, 235, 250)     # #DCEBFA Pastel Blue Fill
BLUE_LINE = RGBColor(59, 130, 246)    # #3B82F6 Blue Border / Text
LAV_BG = RGBColor(233, 225, 247)      # #E9E1F7 Pastel Lavender Fill
LAV_LINE = RGBColor(139, 92, 246)     # #8B5CF6 Lavender Border / Text
PEACH_BG = RGBColor(252, 227, 214)    # #FCE3D6 Pastel Peach Fill
PEACH_LINE = RGBColor(249, 115, 22)   # #F97316 Peach Border / Text
YELLOW_BG = RGBColor(255, 240, 199)   # #FFF0C7 Pastel Yellow Fill
YELLOW_LINE = RGBColor(234, 179, 8)   # #EAB308 Yellow Border / Text
CORAL_BG = RGBColor(254, 226, 226)    # #FEE2E2 Soft Red / Coral Fill
CORAL_LINE = RGBColor(239, 68, 68)    # #EF4444 Coral Border / Text
WHITE = RGBColor(255, 255, 255)

def set_slide_background(slide):
    """Sets clean warm ivory solid background."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = IVORY

def add_header(slide, slide_num, category, title, lede):
    """Standardized top banner with category pill, title and lede."""
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"{slide_num:02d} / 13  •  {category.upper()}"
    p_cat.font.name = "Segoe UI"
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = SLATE

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(23)
    p_title.font.bold = True
    p_title.font.color.rgb = CHARCOAL

    # Lede / Subheading
    lede_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.30), Inches(11.7), Inches(0.35))
    tf_lede = lede_box.text_frame
    tf_lede.word_wrap = True
    tf_lede.margin_left = tf_lede.margin_top = tf_lede.margin_right = tf_lede.margin_bottom = 0
    p_lede = tf_lede.paragraphs[0]
    p_lede.text = lede
    p_lede.font.name = "Segoe UI"
    p_lede.font.size = Pt(11.5)
    p_lede.font.color.rgb = SLATE

def add_footer(slide, slide_num):
    """Subtle bottom footer with branding and student info."""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(11.733), Inches(0.3))
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"NEXUS — AI-Powered Supply Chain Intelligence & Optimization  |  Shreyansh Uttam  |  Slide {slide_num} of 13"
    p.font.name = "Segoe UI"
    p.font.size = Pt(8.5)
    p.font.color.rgb = MUTED

def add_card(slide, left, top, width, height, title=None, fill_color=CREAM, border_color=BORDER_LIGHT, border_width=Pt(1)):
    """Creates a clean rounded card container with pastel styling."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = border_width
    
    if title:
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.32))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
    return shape

def add_speaker_notes(slide, notes_text):
    """Adds concise, spoken-ready speaker notes (30-60s)."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text.strip()

# =============================================================
# SLIDE BUILDERS (1 TO 13)
# =============================================================

def build_slide_01(prs):
    """SLIDE 01 — COVER"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Large container card
    add_card(slide, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), fill_color=CREAM, border_color=BORDER_LIGHT)

    # Left content box
    tb_left = slide.shapes.add_textbox(Inches(1.3), Inches(1.3), Inches(5.6), Inches(4.8))
    tf = tb_left.text_frame
    tf.word_wrap = True

    # Category Pill
    p_pill = tf.paragraphs[0]
    p_pill.text = "PROJECT REVIEW & TECHNICAL EVALUATION"
    p_pill.font.name = "Segoe UI"
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = BLUE_LINE

    # Brand Title
    p_title = tf.add_paragraph()
    p_title.text = "NEXUS"
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(56)
    p_title.font.bold = True
    p_title.font.color.rgb = CHARCOAL
    p_title.space_before = Pt(6)

    # Subtitle
    p_sub = tf.add_paragraph()
    p_sub.text = "AI-Powered Supply Chain Intelligence,\nRisk Prediction & Optimization"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = SLATE
    p_sub.space_before = Pt(4)

    # Small loop pill
    p_loop = tf.add_paragraph()
    p_loop.text = "MONITOR → PREDICT → IMPACT → SIMULATE → OPTIMIZE → RECOMMEND"
    p_loop.font.name = "Segoe UI"
    p_loop.font.size = Pt(9.5)
    p_loop.font.bold = True
    p_loop.font.color.rgb = MINT_LINE
    p_loop.space_before = Pt(14)

    # Candidate details
    p_pres = tf.add_paragraph()
    p_pres.text = "Shreyansh Uttam\nB.Tech CSE (Artificial Intelligence & Machine Learning)\nVIT Bhopal University"
    p_pres.font.name = "Segoe UI"
    p_pres.font.size = Pt(11)
    p_pres.font.color.rgb = CHARCOAL
    p_pres.space_before = Pt(20)

    # Right: Abstract Supply-Chain Network Visual
    cover_img = ASSETS_DIR / "pastel_cover_network.png"
    if cover_img.exists():
        slide.shapes.add_picture(str(cover_img), Inches(7.0), Inches(1.1), Inches(5.2), Inches(5.2))

    add_footer(slide, 1)
    add_speaker_notes(slide, """
Good morning, respected faculty members. My name is Shreyansh Uttam, and today I am presenting NEXUS—an AI-powered supply chain intelligence, risk prediction, and optimization platform.

Modern supply chains are complex multi-echelon networks that face constant volatility. NEXUS unifies empirical data, machine learning forecasting, graph-based impact propagation, what-if simulation, and mathematical optimization into a single closed decision loop. In the next 10 minutes, I will show you how NEXUS detects an upstream disruption, traces its downstream network impact, and automatically computes the optimal reallocation in milliseconds.
""")

def build_slide_02(prs):
    """SLIDE 02 — THE PROBLEM"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 2, "PROBLEM MOTIVATION",
               "Supply chains don't fail at one point.",
               "A small disruption can cascade across the entire network.")

    # 3 Large Pastel Cards
    cards_data = [
        ("FRAGMENTED DATA", "Forecasting, inventory, and logistics exist in separate systems, preventing unified end-to-end visibility.", BLUE_BG, BLUE_LINE),
        ("DELAYED DECISIONS", "Teams often identify the impact only after safety inventory starts running out at distribution hubs.", YELLOW_BG, YELLOW_LINE),
        ("RIGID RESPONSE", "Static supplier allocations make it difficult to quickly reroute supply when a primary facility fails.", PEACH_BG, PEACH_LINE)
    ]

    card_w = Inches(3.7)
    card_h = Inches(3.1)
    start_x = Inches(0.8)
    gap = Inches(0.316)

    for idx, (title, desc, bg, border) in enumerate(cards_data):
        x = start_x + idx * (card_w + gap)
        add_card(slide, x, Inches(1.8), card_w, card_h, fill_color=bg, border_color=border, border_width=Pt(1.5))
        
        tb = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.1), card_w - Inches(0.5), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_num = tf.paragraphs[0]
        p_num.text = f"0{idx+1}"
        p_num.font.name = "Segoe UI"
        p_num.font.size = Pt(28)
        p_num.font.bold = True
        p_num.font.color.rgb = border

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = CHARCOAL
        p_t.space_before = Pt(4)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = CHARCOAL
        p_d.space_before = Pt(8)

    # Bottom Disruption Flow
    add_card(slide, Inches(0.8), Inches(5.15), Inches(11.733), Inches(1.6), fill_color=CREAM, border_color=CORAL_LINE, border_width=Pt(1.2))
    tb_flow = slide.shapes.add_textbox(Inches(1.1), Inches(5.3), Inches(11.133), Inches(1.3))
    tf_f = tb_flow.text_frame
    tf_f.word_wrap = True

    p_fh = tf_f.paragraphs[0]
    p_fh.text = "THE DISRUPTION CASCADE:"
    p_fh.font.name = "Segoe UI"
    p_fh.font.size = Pt(10)
    p_fh.font.bold = True
    p_fh.font.color.rgb = CORAL_LINE

    p_fb = tf_f.add_paragraph()
    p_fb.text = "Supplier Disruption  →  Inventory Depletion  →  Production Starvation  →  Customer Shortages"
    p_fb.font.name = "Segoe UI"
    p_fb.font.size = Pt(14)
    p_fb.font.bold = True
    p_fb.font.color.rgb = CHARCOAL
    p_fb.space_before = Pt(4)

    p_fs = tf_f.add_paragraph()
    p_fs.text = "Without automated intelligence, a regional boiler failure or strike halts downstream retail deliveries within days."
    p_fs.font.name = "Segoe UI"
    p_fs.font.size = Pt(9.5)
    p_fs.font.color.rgb = SLATE
    p_fs.space_before = Pt(4)

    add_footer(slide, 2)
    add_speaker_notes(slide, """
Supply chains do not fail at a single point—they cascade. 

Today's enterprise operations suffer from three major bottlenecks: First, data is fragmented between ERP, WMS, and planning spreadsheets. Second, decisions are delayed because teams only notice stockouts when warehouse shelves are already empty. And third, contracts are rigid—when a primary vendor shuts down, businesses struggle to quickly calculate alternatives.

As shown at the bottom, an upstream supplier failure ripples through warehouses and factories, ultimately leaving retail demand unfulfilled. To solve this, we need a unified system that connects signals to decisions.
""")

def build_slide_03(prs):
    """SLIDE 03 — WHAT IS NEXUS?"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 3, "SYSTEM CONCEPT",
               "One disruption. One decision loop.",
               "A unified decision intelligence platform closing the gap between prediction and operational action.")

    # Loop Visual in center
    loop_img = ASSETS_DIR / "pastel_decision_loop.png"
    if loop_img.exists():
        slide.shapes.add_picture(str(loop_img), Inches(0.8), Inches(1.85), Inches(11.733), Inches(4.3))

    # Bottom Core Principle Pill
    add_card(slide, Inches(1.8), Inches(6.25), Inches(9.733), Inches(0.65), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_bot = slide.shapes.add_textbox(Inches(2.0), Inches(6.32), Inches(9.333), Inches(0.5))
    tf_b = tb_bot.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "Prediction becomes useful only when it leads to a decision."
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = CHARCOAL
    p_b.alignment = PP_ALIGN.CENTER

    add_footer(slide, 3)
    add_speaker_notes(slide, """
This is the core concept of NEXUS: a closed six-stage decision loop.

Many AI projects stop at prediction—they forecast a number or flag a risk, but leave the manager wondering what to do. NEXUS answers six progressive questions: 
MONITOR asks: what is happening right now?
PREDICT asks: what is likely to happen next?
IMPACT asks: what downstream warehouses and products will be affected?
SIMULATE asks: what if we intervene?
OPTIMIZE asks: what is the best feasible allocation of inventory and transit?
And RECOMMEND asks: what exact directives should the executive approve?

Prediction becomes useful only when it directly informs an operational decision.
""")

def build_slide_04(prs):
    """SLIDE 04 — DATA FOUNDATION"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 4, "DATA FOUNDATION",
               "Built on empirical data + a calibrated digital twin",
               "Combining real-world sales volatility and delivery delays with a multi-echelon spatial topology.")

    cards_data = [
        ("WALMART", "Public Empirical Data", "421,570", "Weekly historical records across 45 stores and 81 departments.", "DEMAND FORECASTING", BLUE_BG, BLUE_LINE),
        ("DATACO", "Public Empirical Data", "35,000+", "Representative shipments (180,519 corpus) with authentic delivery delay records.", "SUPPLIER RISK + BEHAVIOR", MINT_BG, MINT_LINE),
        ("INDIAN TWIN", "Calibrated Synthetic", "83 Nodes", "160 Corridors, 50,000 Orders, and 500 SKUs mapped across Indian freight corridors.", "SIMULATION & OPTIMIZATION", LAV_BG, LAV_LINE)
    ]

    card_w = Inches(3.7)
    card_h = Inches(4.3)
    start_x = Inches(0.8)
    gap = Inches(0.316)

    for idx, (title, nature, stat, desc, purpose, bg, border) in enumerate(cards_data):
        x = start_x + idx * (card_w + gap)
        add_card(slide, x, Inches(1.8), card_w, card_h, fill_color=bg, border_color=border, border_width=Pt(1.5))
        
        tb = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.05), card_w - Inches(0.5), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_nat = tf.paragraphs[0]
        p_nat.text = nature.upper()
        p_nat.font.name = "Segoe UI"
        p_nat.font.size = Pt(9)
        p_nat.font.bold = True
        p_nat.font.color.rgb = border

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = CHARCOAL
        p_t.space_before = Pt(2)

        p_stat = tf.add_paragraph()
        p_stat.text = stat
        p_stat.font.name = "Segoe UI"
        p_stat.font.size = Pt(26)
        p_stat.font.bold = True
        p_stat.font.color.rgb = CHARCOAL
        p_stat.space_before = Pt(8)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = SLATE
        p_desc.space_before = Pt(4)

        p_purp_lbl = tf.add_paragraph()
        p_purp_lbl.text = "PRIMARY PURPOSE:"
        p_purp_lbl.font.name = "Segoe UI"
        p_purp_lbl.font.size = Pt(8.5)
        p_purp_lbl.font.bold = True
        p_purp_lbl.font.color.rgb = border
        p_purp_lbl.space_before = Pt(12)

        p_purp = tf.add_paragraph()
        p_purp.text = purpose
        p_purp.font.name = "Segoe UI"
        p_purp.font.size = Pt(11)
        p_purp.font.bold = True
        p_purp.font.color.rgb = CHARCOAL
        p_purp.space_before = Pt(2)

    # Academic Integrity Note
    tb_note = slide.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.5))
    tf_n = tb_note.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = "Academic Integrity: Empirical datasets ground the models; the Indian network is a simulated digital twin."
    p_n.font.name = "Segoe UI"
    p_n.font.size = Pt(9.5)
    p_n.font.italic = True
    p_n.font.color.rgb = MUTED
    p_n.alignment = PP_ALIGN.CENTER

    add_footer(slide, 4)
    add_speaker_notes(slide, """
Every engineering system must stand on credible data. In supply chain management, real enterprise operational data is proprietary and confidential. Therefore, NEXUS employs a hybrid data strategy.

For demand forecasting, we use Walmart's empirical dataset of 421,570 weekly sales records from Kaggle, capturing real retail seasonality and holiday surges. For supplier risk, we utilize the DataCo Global Supply Chain dataset from Mendeley, containing over 35,000 shipments with real delivery delays. 

We then map these statistical distributions onto an 83-node Indian logistics network. As clearly stated at the bottom, this Indian topology is a calibrated synthetic digital twin, ensuring academic honesty while maintaining realistic multi-echelon complexity.
""")

def build_slide_05(prs):
    """SLIDE 05 — DIGITAL SUPPLY CHAIN TWIN"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 5, "DIGITAL TWIN TOPOLOGY",
               "A supply chain you can see, trace and stress-test.",
               "An 83-node, 160-corridor digital representation across five operational tiers.")

    # Left: Visual Network Diagram
    net_img = ASSETS_DIR / "pastel_supply_chain_twin.png"
    if net_img.exists():
        slide.shapes.add_picture(str(net_img), Inches(0.8), Inches(1.8), Inches(8.0), Inches(4.9))

    # Right: KPI Stack Cards
    kpis = [
        ("83", "Facility Nodes", "Suppliers, plants, warehouses, hubs, zones", BLUE_BG, BLUE_LINE),
        ("160", "Freight Corridors", "Multimodal highway & rail freight lanes", MINT_BG, MINT_LINE),
        ("30", "Demand Zones", "Tier-1 & Tier-2 consumption centers", LAV_BG, LAV_LINE),
        ("₹45.9M", "Inventory Value", "500 SKUs tracked across central warehouses", YELLOW_BG, YELLOW_LINE)
    ]

    kpi_w = Inches(3.4)
    kpi_h = Inches(1.1)
    gap_y = Inches(0.16)
    start_y = Inches(1.8)

    for idx, (num, label, desc, bg, border) in enumerate(kpis):
        y = start_y + idx * (kpi_h + gap_y)
        add_card(slide, Inches(9.133), y, kpi_w, kpi_h, fill_color=bg, border_color=border, border_width=Pt(1.2))
        
        tb = slide.shapes.add_textbox(Inches(9.3), y + Inches(0.1), kpi_w - Inches(0.35), kpi_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = num
        p_num.font.name = "Segoe UI"
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = CHARCOAL

        p_lbl = tf.add_paragraph()
        p_lbl.text = f"{label} — {desc}"
        p_lbl.font.name = "Segoe UI"
        p_lbl.font.size = Pt(8.5)
        p_lbl.font.color.rgb = SLATE

    add_footer(slide, 5)
    add_speaker_notes(slide, """
Here is the NEXUS Digital Twin. It models a complete multi-echelon network spanning 83 facilities and 160 multimodal corridors across India.

The network is partitioned into five distinct echelons: 20 suppliers, 8 manufacturing plants, 10 central distribution warehouses tracking 45.9 million rupees in inventory, 15 transit hubs, and 30 consumer demand zones. 

On the map visual, you can see our highlighted disruption test: an 80% failure at supplier SUP_001 in Pune cascading down to Western Warehouse WH_01, alongside the green rerouting corridor from alternative vendor SUP_005 in Jamshedpur. The digital twin makes every dependency visible and computable.
""")

def build_slide_06(prs):
    """SLIDE 06 — INTELLIGENCE LAYER"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 6, "INTELLIGENCE LAYER",
               "Three models. Three questions.",
               "Machine learning detects the operational signal before failure cascades.")

    models_data = [
        ("01 — DEMAND FORECASTING", "XGBoost Regressor", '"How much will we need?"', "4-week multi-step demand forecast with seasonal confidence intervals.", BLUE_BG, BLUE_LINE),
        ("02 — SUPPLIER RISK", "XGBoost Classifier", '"Who is likely to fail?"', "Vendor disruption probability based on lead-time volatility and delay history.", MINT_BG, MINT_LINE),
        ("03 — ANOMALY DETECTION", "Isolation Forest", '"What looks unusual?"', "Real-time unsupervised detection of sudden demand surges and transit outliers.", LAV_BG, LAV_LINE)
    ]

    card_w = Inches(3.7)
    card_h = Inches(4.1)
    start_x = Inches(0.8)
    gap = Inches(0.316)

    for idx, (title, arch, question, output_desc, bg, border) in enumerate(models_data):
        x = start_x + idx * (card_w + gap)
        add_card(slide, x, Inches(1.8), card_w, card_h, fill_color=bg, border_color=border, border_width=Pt(1.5))
        
        tb = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.05), card_w - Inches(0.5), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = border

        p_arch = tf.add_paragraph()
        p_arch.text = arch
        p_arch.font.name = "Segoe UI"
        p_arch.font.size = Pt(18)
        p_arch.font.bold = True
        p_arch.font.color.rgb = CHARCOAL
        p_arch.space_before = Pt(4)

        p_q_lbl = tf.add_paragraph()
        p_q_lbl.text = "CORE QUESTION:"
        p_q_lbl.font.name = "Segoe UI"
        p_q_lbl.font.size = Pt(8.5)
        p_q_lbl.font.bold = True
        p_q_lbl.font.color.rgb = border
        p_q_lbl.space_before = Pt(14)

        p_q = tf.add_paragraph()
        p_q.text = question
        p_q.font.name = "Segoe UI"
        p_q.font.size = Pt(13)
        p_q.font.bold = True
        p_q.font.color.rgb = CHARCOAL
        p_q.space_before = Pt(2)

        p_out_lbl = tf.add_paragraph()
        p_out_lbl.text = "OUTPUT SIGNAL:"
        p_out_lbl.font.name = "Segoe UI"
        p_out_lbl.font.size = Pt(8.5)
        p_out_lbl.font.bold = True
        p_out_lbl.font.color.rgb = border
        p_out_lbl.space_before = Pt(14)

        p_out = tf.add_paragraph()
        p_out.text = output_desc
        p_out.font.name = "Segoe UI"
        p_out.font.size = Pt(9.5)
        p_out.font.color.rgb = SLATE
        p_out.space_before = Pt(2)

    # Bottom statement pill
    add_card(slide, Inches(1.8), Inches(6.15), Inches(9.733), Inches(0.65), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_bot = slide.shapes.add_textbox(Inches(2.0), Inches(6.22), Inches(9.333), Inches(0.5))
    tf_b = tb_bot.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "ML detects the signal. NEXUS then traces its operational impact."
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = CHARCOAL
    p_b.alignment = PP_ALIGN.CENTER

    add_footer(slide, 6)
    add_speaker_notes(slide, """
The Intelligence Layer is built around three specific operational questions rather than generic AI buzzwords.

First, Demand Forecasting asks: 'How much will we need?' It uses an XGBoost Regressor trained on 14 lag and calendar features to output a 4-week demand curve.
Second, Supplier Risk asks: 'Who is likely to fail?' An XGBoost Classifier predicts vendor failure probability from lead-time volatility and delay frequency.
Third, Anomaly Detection asks: 'What looks unusual right now?' An unsupervised Isolation Forest flags unprecedented order surges or highway transit delays.

Critically, machine learning provides the detection signal; NEXUS then determines what that signal actually means for the network.
""")

def build_slide_07(prs):
    """SLIDE 07 — MODEL COMPARISON"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 7, "EMPIRICAL BENCHMARKING",
               "Does the model actually improve the decision signal?",
               "Empirical out-of-sample evaluation against established industry baselines.")

    # Left Card: Demand Forecasting
    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_fc = slide.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(5.1), Inches(1.0))
    tf_fc = tb_fc.text_frame
    p_fc_h = tf_fc.paragraphs[0]
    p_fc_h.text = "DEMAND FORECASTING (WALMART DATASET)"
    p_fc_h.font.name = "Segoe UI"
    p_fc_h.font.size = Pt(11)
    p_fc_h.font.bold = True
    p_fc_h.font.color.rgb = CHARCOAL

    # Table Left
    headers_fc = ["Model Architecture", "RMSE", "sMAPE"]
    rows_fc = [
        ["Naive Persistence", "136,255.10", "7.20%"],
        ["4-Week Moving Average", "153,446.84", "8.42%"],
        ["XGBoost (NEXUS Model)", "96,042.50", "4.35%"]
    ]
    create_table_pastel(slide, Inches(1.05), Inches(2.55), Inches(5.1), Inches(1.8), headers_fc, rows_fc, highlight_row=2, highlight_bg=MINT_BG, highlight_color=MINT_LINE)

    # Callout pill Left
    add_card(slide, Inches(1.05), Inches(4.8), Inches(5.1), Inches(1.5), fill_color=MINT_BG, border_color=MINT_LINE)
    tb_cp1 = slide.shapes.add_textbox(Inches(1.2), Inches(4.95), Inches(4.8), Inches(1.2))
    tf_cp1 = tb_cp1.text_frame
    p1 = tf_cp1.paragraphs[0]
    p1.text = "29.51% RMSE Reduction"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = CHARCOAL
    p2 = tf_cp1.add_paragraph()
    p2.text = "XGBoost lowers forecasting error by 29.51% over Naive persistence and 37.41% over 4-Week Moving Average, significantly dampening the Bullwhip effect."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = SLATE
    p2.space_before = Pt(3)

    # Right Card: Supplier Risk
    add_card(slide, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_rk = slide.shapes.add_textbox(Inches(7.05), Inches(2.0), Inches(5.2), Inches(1.0))
    tf_rk = tb_rk.text_frame
    p_rk_h = tf_rk.paragraphs[0]
    p_rk_h.text = "SUPPLIER RISK (DATACO / 10 UNSEEN VENDORS)"
    p_rk_h.font.name = "Segoe UI"
    p_rk_h.font.size = Pt(11)
    p_rk_h.font.bold = True
    p_rk_h.font.color.rgb = CHARCOAL

    headers_rk = ["Model Architecture", "Precision", "Recall", "F1"]
    rows_rk = [
        ["Logistic Regression", "0.7531", "0.7349", "0.7439"],
        ["XGBoost (NEXUS Model)", "0.7442", "0.7711", "0.7574"]
    ]
    create_table_pastel(slide, Inches(7.05), Inches(2.55), Inches(5.2), Inches(1.5), headers_rk, rows_rk, highlight_row=1, highlight_bg=BLUE_BG, highlight_color=BLUE_LINE)

    add_card(slide, Inches(7.05), Inches(4.8), Inches(5.2), Inches(1.5), fill_color=BLUE_BG, border_color=BLUE_LINE)
    tb_cp2 = slide.shapes.add_textbox(Inches(7.2), Inches(4.95), Inches(4.9), Inches(1.2))
    tf_cp2 = tb_cp2.text_frame
    p3 = tf_cp2.paragraphs[0]
    p3.text = "77.11% Out-of-Sample Recall"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = CHARCOAL
    p4 = tf_cp2.add_paragraph()
    p4.text = "Evaluated via GroupShuffleSplit on 10 completely unseen vendors. High recall ensures early warning of supply failure where missing a disruption carries massive penalty costs."
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(8.5)
    p4.font.color.rgb = SLATE
    p4.space_before = Pt(3)

    add_footer(slide, 7)
    add_speaker_notes(slide, """
Here we evaluate whether machine learning actually improves the decision signal compared to established baselines.

On the left, for demand forecasting, our XGBoost Regressor achieves an RMSE of 96,042 and an sMAPE of 4.35%. This is a verified 29.51% reduction in root mean squared error over the Naive baseline, directly dampening Bullwhip order distortion.

On the right, for supplier risk, we evaluated the models across 10 completely unseen test vendors using GroupShuffleSplit. XGBoost achieves an F1-score of 0.7574 with a high recall of 77.11%. In supply chains, high recall is vital because missing an impending supplier collapse is far more catastrophic than a minor false alarm.
""")

def build_slide_08(prs):
    """SLIDE 08 — DECISION MATRIX"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 8, "DECISION MATRIX",
               "From signal → impact → action",
               "Prediction alone does not make decisions; NEXUS connects signals to operational responses.")

    # Table Card
    add_card(slide, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.0), fill_color=CREAM, border_color=BORDER_LIGHT)

    headers_dm = ["Signal Detected", "Operational Meaning", "NEXUS Prescriptive Response"]
    rows_dm = [
        ["● Demand surge", "Safety stock may deplete prematurely", "Dynamically increase procurement allocations"],
        ["● Supplier capacity loss", "Critical component shortage risk", "Shift sourcing to qualified secondary vendors"],
        ["● Route delay / blockage", "Transit lead-time increases by 2-5 days", "Reroute freight via multimodal rail/road"],
        ["● Low inventory runway", "Imminent warehouse stockout (< 7 days)", "Trigger inter-warehouse safety stock transfer"],
        ["● Multi-node shock", "Multiple constraints interact across tiers", "Execute Google OR-Tools global optimization"]
    ]

    create_table_pastel(slide, Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.4), headers_dm, rows_dm,
                        col_widths=[Inches(3.3), Inches(4.0), Inches(4.033)])

    # Bottom Takeaway Card
    add_card(slide, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.75), fill_color=MINT_BG, border_color=MINT_LINE)
    tb_b = slide.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.333), Inches(0.55))
    tf_b = tb_b.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "NEXUS does not stop at prediction. It connects the prediction to an operational response."
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = CHARCOAL
    p_b.alignment = PP_ALIGN.CENTER

    add_footer(slide, 8)
    add_speaker_notes(slide, """
This slide shows how NEXUS bridges the gap between signals and actions through our Decision Matrix.

Each detected signal has a specific physical consequence and a prescribed operational response. For example, when a supplier loses capacity, the system doesn't just sound an alarm—it identifies component risk and initiates secondary vendor sourcing. When a highway corridor is blocked, it routes freight through rail corridors. And when multiple shocks strike simultaneously, it triggers global optimization.

This answers a fundamental faculty question: NEXUS is not just an analytics dashboard—it is an automated decision-support engine.
""")

def build_slide_09(prs):
    """SLIDE 09 — DISRUPTION SIMULATION"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 9, "DISRUPTION SIMULATION",
               "What happens when a critical supplier loses capacity?",
               "Simulating an 80% capacity shock on supplier SUP_001 in Pune.")

    # Left: Disruption Shock & Cascade Story (3 vertical stacked cards)
    steps = [
        ("SHOCK", "SUP_001 (Tata AutoComp Pune)\n80% Capacity Loss for 10 Days", CORAL_BG, CORAL_LINE),
        ("CASCADE", "Western Warehouse (WH_01)\n1,420 units stock vs 280 units/day burn rate = 5.07 Days Runway", YELLOW_BG, YELLOW_LINE),
        ("IMPACT", "69,843 Units Potential Shortage\nService level drops to 54.57% without intervention", PEACH_BG, PEACH_LINE)
    ]

    card_w = Inches(5.6)
    card_h = Inches(1.4)
    start_y = Inches(1.8)
    gap_y = Inches(0.2)

    for idx, (lbl, text, bg, border) in enumerate(steps):
        y = start_y + idx * (card_h + gap_y)
        add_card(slide, Inches(0.8), y, card_w, card_h, fill_color=bg, border_color=border, border_width=Pt(1.2))
        
        tb = slide.shapes.add_textbox(Inches(1.05), y + Inches(0.15), card_w - Inches(0.5), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_lbl = tf.paragraphs[0]
        p_lbl.text = lbl
        p_lbl.font.name = "Segoe UI"
        p_lbl.font.size = Pt(10)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = border

        p_t = tf.add_paragraph()
        p_t.text = text
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(9.5)
        p_t.font.bold = True
        p_t.font.color.rgb = CHARCOAL
        p_t.space_before = Pt(2)

    # Right: The Key Insight Card
    add_card(slide, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.6), fill_color=CREAM, border_color=MINT_LINE, border_width=Pt(1.5))
    tb_ins = slide.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.133), Inches(4.0))
    tf_ins = tb_ins.text_frame
    tf_ins.word_wrap = True

    p_q = tf_ins.paragraphs[0]
    p_q.text = "THE CENTRAL STRATEGIC QUESTION:"
    p_q.font.name = "Segoe UI"
    p_q.font.size = Pt(10)
    p_q.font.bold = True
    p_q.font.color.rgb = SLATE

    p_qh = tf_ins.add_paragraph()
    p_qh.text = '"Does the network actually lack capacity?"'
    p_qh.font.name = "Segoe UI"
    p_qh.font.size = Pt(18)
    p_qh.font.bold = True
    p_qh.font.color.rgb = CHARCOAL
    p_qh.space_before = Pt(4)

    p_ans = tf_ins.add_paragraph()
    p_ans.text = "+160,750 Units"
    p_ans.font.name = "Segoe UI"
    p_ans.font.size = Pt(36)
    p_ans.font.bold = True
    p_ans.font.color.rgb = MINT_LINE
    p_ans.space_before = Pt(14)

    p_sub = tf_ins.add_paragraph()
    p_sub.text = "Net Network Capacity Surplus Post-Shock"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(11)
    p_sub.font.bold = True
    p_sub.font.color.rgb = CHARCOAL
    p_sub.space_before = Pt(2)

    p_core = tf_ins.add_paragraph()
    p_core.text = "Total Network Demand = 153,600 units | Remaining Supplier Capacity = 314,350 units."
    p_core.font.name = "Segoe UI"
    p_core.font.size = Pt(9.5)
    p_core.font.color.rgb = SLATE
    p_core.space_before = Pt(6)

    p_takeaway = tf_ins.add_paragraph()
    p_takeaway.text = "Key Takeaway: The problem is allocation — not total network capacity."
    p_takeaway.font.name = "Segoe UI"
    p_takeaway.font.size = Pt(11.5)
    p_takeaway.font.bold = True
    p_takeaway.font.color.rgb = CHARCOAL
    p_takeaway.space_before = Pt(12)

    add_footer(slide, 9)
    add_speaker_notes(slide, """
To evaluate NEXUS in action, we simulate a severe operational crisis: supplier SUP_001 in Pune loses 80% capacity for 10 days. 

On the left, you see the cascade: Western Warehouse WH_01 only has 5.07 days of stock runway remaining, meaning complete stockout on Day 6. Without reallocation, 69,843 units of customer demand would be dropped, plunging service levels to 54.57%.

However, look at the right side: does the network actually lack capacity? No! Even after the shock, remaining suppliers possess 314,000 units of capacity against a network demand of 153,600 units—a net surplus of 160,000 units. The crisis exists solely because traditional allocations are rigid. This brings us directly to optimization.
""")

def build_slide_10(prs):
    """SLIDE 10 — OPTIMIZATION RESULT"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 10, "OPTIMIZATION OUTCOME",
               "NEXUS turns the crisis into an allocation decision.",
               "Google OR-Tools multi-echelon linear optimization achieves global cost minimization in 0.0094 seconds.")

    # High-Res Pastel Optimization Graphic in center
    opt_img = ASSETS_DIR / "pastel_optimization_diff.png"
    if opt_img.exists():
        slide.shapes.add_picture(str(opt_img), Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.3))

    # Bottom Academic Note Card
    add_card(slide, Inches(1.5), Inches(6.25), Inches(10.333), Inches(0.65), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_bot = slide.shapes.add_textbox(Inches(1.7), Inches(6.32), Inches(9.933), Inches(0.5))
    tf_b = tb_bot.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "Simulated stress scenario using Google OR-Tools multi-echelon optimization (Solver runtime: 0.0094s)."
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(10)
    p_b.font.italic = True
    p_b.font.color.rgb = SLATE
    p_b.alignment = PP_ALIGN.CENTER

    add_footer(slide, 10)
    add_speaker_notes(slide, """
This is our flagship result: baseline operations versus NEXUS optimization.

Under unoptimized baseline operations, the company incurs 41.22 million rupees in total operational cost—dominated by 24.45 million in contractual stockout penalties—with a service level of just 54.57%.

NEXUS changes this completely. In 9.4 milliseconds, Google OR-Tools computes a global multi-echelon reallocation. By proactively spending 5 million rupees in secondary procurement and freight rerouting from Jamshedpur and Pantnagar, it completely eliminates all 24.45 million rupees of stockout penalties. 

The result is a 53.19% reduction in total operational cost—saving nearly 22 million rupees—while restoring customer fulfillment to 100%.
""")

def build_slide_11(prs):
    """SLIDE 11 — THE PRODUCT"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 11, "CONTROL ROOM INTERFACE",
               "From model pipeline to working control room",
               "An interactive enterprise SaaS control room built with React 19, TypeScript, and FastAPI.")

    # Visual collage in center
    prod_img = ASSETS_DIR / "pastel_product_collage.png"
    if prod_img.exists():
        slide.shapes.add_picture(str(prod_img), Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.2))

    # Bottom Tech Stack Banner
    add_card(slide, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.75), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_tech = slide.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.333), Inches(0.55))
    tf_t = tb_tech.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "TECH STACK: React 19  •  TypeScript  •  FastAPI  •  Leaflet GIS  •  Recharts  •  Google OR-Tools  •  SQLite / PostgreSQL"
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(10.5)
    p_t.font.bold = True
    p_t.font.color.rgb = CHARCOAL
    p_t.alignment = PP_ALIGN.CENTER

    add_footer(slide, 11)
    add_speaker_notes(slide, """
A decision algorithm is only useful if managers can interact with it. Here you see the live NEXUS control room.

Built with React 19, TypeScript, and Tailwind CSS, the platform organizes intelligence into five dedicated views: MONITOR on the executive dashboard, TRACE on the interactive Leaflet GIS map, PREDICT on demand charts, OPTIMIZE on cost waterfall comparisons, and DECIDE on executive action cards.

The frontend is dynamically code-split into lightweight bundles of 365 kilobytes and builds in under 2.5 seconds, delivering a snappy, enterprise-grade user experience.
""")

def build_slide_12(prs):
    """SLIDE 12 — VALIDATION"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 12, "QUALITY ASSURANCE",
               "Built, tested and verified.",
               "Comprehensive full-stack verification ensuring complete system stability and zero silent failures.")

    # 4 Large KPI Cards
    kpis_val = [
        ("29 / 29", "Backend Tests Passed", "Pytest suite (2.66s) covering ML, graph & LP", MINT_BG, MINT_LINE),
        ("10 / 10", "Frontend Tests Passed", "Vitest unit tests for UI and currency formats", BLUE_BG, BLUE_LINE),
        ("0", "Build Errors", "TypeScript strict compiler & Vite 8.3 bundler", LAV_BG, LAV_LINE),
        ("20", "REST API Endpoints", "FastAPI endpoints verified with HTTP 200 OK", YELLOW_BG, YELLOW_LINE)
    ]

    card_w = Inches(2.75)
    card_h = Inches(2.8)
    start_x = Inches(0.8)
    gap = Inches(0.244)

    for idx, (stat, label, desc, bg, border) in enumerate(kpis_val):
        x = start_x + idx * (card_w + gap)
        add_card(slide, x, Inches(1.8), card_w, card_h, fill_color=bg, border_color=border, border_width=Pt(1.5))
        
        tb = slide.shapes.add_textbox(x + Inches(0.2), Inches(2.1), card_w - Inches(0.4), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_s = tf.paragraphs[0]
        p_s.text = stat
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(36)
        p_s.font.bold = True
        p_s.font.color.rgb = CHARCOAL

        p_l = tf.add_paragraph()
        p_l.text = label
        p_l.font.name = "Segoe UI"
        p_l.font.size = Pt(11)
        p_l.font.bold = True
        p_l.font.color.rgb = CHARCOAL
        p_l.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = SLATE
        p_d.space_before = Pt(4)

    # Architecture strip below
    add_card(slide, Inches(0.8), Inches(4.85), Inches(11.733), Inches(1.1), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_strip = slide.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.333), Inches(0.8))
    tf_st = tb_strip.text_frame
    p_st1 = tf_st.paragraphs[0]
    p_st1.text = "INTEGRATION PIPELINE:  React UI  →  FastAPI  →  ML / Graph / Simulation  →  Google OR-Tools  →  SQLite Persistence"
    p_st1.font.name = "Segoe UI"
    p_st1.font.size = Pt(10)
    p_st1.font.bold = True
    p_st1.font.color.rgb = CHARCOAL
    p_st1.alignment = PP_ALIGN.CENTER

    p_st2 = tf_st.add_paragraph()
    p_st2.text = "Fixed random seeds (seed=42)  •  Zero silent mock fallbacks  •  100% reproducible execution"
    p_st2.font.name = "Segoe UI"
    p_st2.font.size = Pt(9)
    p_st2.font.color.rgb = SLATE
    p_st2.space_before = Pt(3)
    p_st2.alignment = PP_ALIGN.CENTER

    # Academic disclaimer footer
    tb_disc = slide.shapes.add_textbox(Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.5))
    tf_dc = tb_disc.text_frame
    p_dc = tf_dc.paragraphs[0]
    p_dc.text = "Academic Honesty: The platform does not claim live connection to SAP/Oracle or real GPS telematics; ERP dispatch is simulated."
    p_dc.font.name = "Segoe UI"
    p_dc.font.size = Pt(9)
    p_dc.font.italic = True
    p_dc.font.color.rgb = MUTED
    p_dc.alignment = PP_ALIGN.CENTER

    add_footer(slide, 12)
    add_speaker_notes(slide, """
Every engineering system must be verified. NEXUS has undergone rigorous automated testing across both backend and frontend.

On the backend, all 29 pytest tests pass with 100% success in 2.66 seconds, validating schemas, forecasting algorithms, graph traversal, and optimization constraints. On the frontend, all 10 Vitest tests pass, and the production build compiles with zero TypeScript errors. 

Furthermore, our codebase adheres to strict credibility standards: there are zero silent mock fallbacks in the production frontend. Every figure traces directly to the live backend. And as noted at the bottom, we transparently state that live enterprise ERP and GPS connections are simulated.
""")

def build_slide_13(prs):
    """SLIDE 13 — CLOSING"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, 13, "CONCLUSION & DEFENSE",
               "NEXUS in one picture.",
               "Transforming fragmented supply chain data into mathematically verified, actionable decisions.")

    # High-level Journey Banner
    add_card(slide, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.85), fill_color=CREAM, border_color=BLUE_LINE, border_width=Pt(1.2))
    tb_j = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(0.5))
    tf_j = tb_j.text_frame
    p_j = tf_j.paragraphs[0]
    p_j.text = "DATA   →   UNDERSTAND   →   PREDICT   →   TRACE   →   SIMULATE   →   OPTIMIZE   →   DECIDE"
    p_j.font.name = "Segoe UI"
    p_j.font.size = Pt(13)
    p_j.font.bold = True
    p_j.font.color.rgb = CHARCOAL
    p_j.alignment = PP_ALIGN.CENTER

    # Core Summary Card
    add_card(slide, Inches(0.8), Inches(2.9), Inches(11.733), Inches(1.6), fill_color=MINT_BG, border_color=MINT_LINE)
    tb_core = slide.shapes.add_textbox(Inches(1.1), Inches(3.1), Inches(11.133), Inches(1.2))
    tf_c = tb_core.text_frame
    tf_c.word_wrap = True
    p_ct = tf_c.paragraphs[0]
    p_ct.text = '"From fragmented supply-chain data to actionable decisions."'
    p_ct.font.name = "Segoe UI"
    p_ct.font.size = Pt(16)
    p_ct.font.bold = True
    p_ct.font.color.rgb = CHARCOAL

    p_cd = tf_c.add_paragraph()
    p_cd.text = "NEXUS bridges the gap between predictive AI and operations research, providing supply chain operators with the tools to foresee disruption, trace cascading consequences, and execute optimal mitigation."
    p_cd.font.name = "Segoe UI"
    p_cd.font.size = Pt(10.5)
    p_cd.font.color.rgb = SLATE
    p_cd.space_before = Pt(4)

    # Closing Presenter & Thank You Card
    add_card(slide, Inches(0.8), Inches(4.75), Inches(11.733), Inches(2.0), fill_color=CREAM, border_color=BORDER_LIGHT)
    tb_pres = slide.shapes.add_textbox(Inches(1.1), Inches(4.95), Inches(11.133), Inches(1.6))
    tf_p = tb_pres.text_frame
    
    p_ty = tf_p.paragraphs[0]
    p_ty.text = "Thank You!   |   Questions & Discussion"
    p_ty.font.name = "Segoe UI"
    p_ty.font.size = Pt(20)
    p_ty.font.bold = True
    p_ty.font.color.rgb = CHARCOAL

    p_dt = tf_p.add_paragraph()
    p_dt.text = "Shreyansh Uttam  •  B.Tech Computer Science & Engineering (AI & ML)\nSchool of Computing Science and Engineering (SCSE)  •  VIT Bhopal University"
    p_dt.font.name = "Segoe UI"
    p_dt.font.size = Pt(11)
    p_dt.font.color.rgb = SLATE
    p_dt.space_before = Pt(6)

    add_footer(slide, 13)
    add_speaker_notes(slide, """
To conclude, NEXUS transforms supply chain management from reactive firefighting into proactive decision intelligence.

By connecting empirical data, digital twin graph analysis, machine learning, scenario simulation, and Google OR-Tools optimization, the platform doesn't just predict problems—it solves them. In our verified demonstration, NEXUS eliminated 24.45 million rupees in stockout penalties, maintained 100% service level, and reduced total operational costs by 53.19% in under 10 milliseconds.

Thank you very much for your time and guidance. I would be delighted to answer any questions and demonstrate the live platform.
""")

# =============================================================
# TABLE UTILITIES (PASTEL THEMED)
# =============================================================

def create_table_pastel(slide, left, top, width, height, headers, rows, col_widths=None, highlight_row=-1, highlight_bg=MINT_BG, highlight_color=MINT_LINE):
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    if col_widths and len(col_widths) == num_cols:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = w

    # Headers
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CREAM
        cell.text_frame.word_wrap = True
        cell.text_frame.margin_left = cell.text_frame.margin_right = Inches(0.08)
        cell.text_frame.margin_top = cell.text_frame.margin_bottom = Inches(0.04)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = SLATE

    # Rows
    for r_idx, row in enumerate(rows):
        is_highlight = (r_idx == highlight_row)
        bg = highlight_bg if is_highlight else (CREAM if r_idx % 2 == 0 else IVORY)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.text_frame.word_wrap = True
            cell.text_frame.margin_left = cell.text_frame.margin_right = Inches(0.08)
            cell.text_frame.margin_top = cell.text_frame.margin_bottom = Inches(0.04)
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.name = "Segoe UI"
            p.font.size = Pt(9)
            if is_highlight:
                p.font.bold = True
                p.font.color.rgb = highlight_color if c_idx > 0 else CHARCOAL
            else:
                p.font.color.rgb = CHARCOAL if c_idx == 0 else SLATE
    return table_shape

# =============================================================
# MAIN ORCHESTRATION & PDF EXPORT
# =============================================================

def generate_pastel_presentation():
    print("=" * 70)
    print(" GENERATING NEXUS 13-SLIDE PASTEL PRESENTATION ")
    print("=" * 70)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    builders = [
        build_slide_01, build_slide_02, build_slide_03, build_slide_04,
        build_slide_05, build_slide_06, build_slide_07, build_slide_08,
        build_slide_09, build_slide_10, build_slide_11, build_slide_12,
        build_slide_13
    ]

    for idx, builder in enumerate(builders, 1):
        print(f"Building Slide {idx:02d} / 13: {builder.__doc__}...")
        builder(prs)

    prs.save(str(PPTX_OUTPUT))
    print(f"\nSuccessfully generated PowerPoint presentation:")
    print(f"-> {PPTX_OUTPUT} ({os.path.getsize(PPTX_OUTPUT):,} bytes)")

    # Save source copy
    import shutil
    shutil.copy2(__file__, SOURCE_COPY)
    print(f"-> Source copy saved to: {SOURCE_COPY}")

    # Export to PDF via PowerPoint COM automation
    export_to_pdf()

def export_to_pdf():
    print("\nAttempting native PDF export via PowerPoint COM automation...")
    ps_cmd = f"""
    $ppt = New-Object -ComObject PowerPoint.Application
    $ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoFalse
    $presentation = $ppt.Presentations.Open('{PPTX_OUTPUT}')
    $presentation.SaveAs('{PDF_OUTPUT}', 32)
    $presentation.Close()
    $ppt.Quit()
    Write-Host "PDF Export Successful: {PDF_OUTPUT}"
    """
    try:
        res = subprocess.run(["powershell", "-Command", ps_cmd], capture_output=True, text=True, timeout=60)
        if res.returncode == 0 and PDF_OUTPUT.exists():
            print(f"-> PDF Successfully generated: {PDF_OUTPUT} ({os.path.getsize(PDF_OUTPUT):,} bytes)")
        else:
            print("PowerPoint COM PDF export message:", res.stdout, res.stderr)
    except Exception as e:
        print("PDF export failed or timed out:", e)

if __name__ == "__main__":
    generate_pastel_presentation()
